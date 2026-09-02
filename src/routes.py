from datetime import datetime, timezone, timedelta
from functools import wraps
from io import BytesIO, StringIO
import csv
import hashlib
import os
import re
import secrets
import uuid
import random
from flask import Blueprint, current_app, jsonify, request, send_file, send_from_directory
from flask_mail import Message
from werkzeug.utils import secure_filename
from flask_jwt_extended import create_access_token, get_jwt, get_jwt_identity, jwt_required
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func, or_
from sqlalchemy.orm import joinedload, selectinload
from . import db, limiter, mail
from .redis_service import set_json, get_json, acquire_lock, release_lock
from .models import Institute, User, AccountToken, StudentProfile, StudentCategory, Course, Batch, Exam, ExamAssignment, Question, QuestionOption, ExamAttempt, AttemptAnswer, AttemptEvent, AuditLog, QuestionBankItem, now

api=Blueprint("api",__name__)
def ok(data=None,message=None,status=200): return jsonify({"success":True,"data":data,"message":message}),status
def fail(message,status=400): return jsonify({"success":False,"message":message}),status
def role_required(role):
    def outer(fn):
        @wraps(fn)
        @jwt_required()
        def inner(*a,**kw):
            if get_jwt().get("role")!=role:
                user=db.session.get(User,uid());db.session.add(AuditLog(institute_id=user.institute_id if user else None,actor_id=user.id if user else None,action="PERMISSION_DENIED",category="SECURITY",severity="WARNING",entity_type="ROUTE",details={"required_role":role,"path":request.path},ip_address=client_ip(),user_agent=request.user_agent.string[:500],request_id=request.headers.get("X-Request-ID") or str(uuid.uuid4())));db.session.commit();return fail("Forbidden",403)
            return fn(*a,**kw)
        return inner
    return outer
def uid(): return int(get_jwt_identity())
def tenant_id():
    claim=get_jwt().get("institute_id")
    return int(claim or db.session.get(User,uid()).institute_id)
def dt(v): return datetime.fromisoformat(v.replace("Z","+00:00")) if v else None
def client_ip():return (request.headers.get("X-Forwarded-For",request.remote_addr or "").split(",")[0].strip())[:64]
def sync(exam):
    t=now()
    if exam.status=="PUBLISHED" and exam.end_at and exam.end_at.replace(tzinfo=timezone.utc) <= t:
        exam.status="COMPLETED"; exam.completed_at=t; db.session.commit()
    return exam
def exam_json(e,detail=False,include_answers=False):
    sync(e); assigned_ids=[assignment.student_id for assignment in e.assignments]; target=e.category.name if e.category else f"{len(assigned_ids)} selected student{'s' if len(assigned_ids)!=1 else ''}"
    d={"id":e.id,"access_code":e.access_code,"title":e.title,"description":e.description,"exam_category":e.exam_category,"target_category":target,"audience_mode":"category" if e.student_category_id else "students","student_category_id":e.student_category_id,"duration":e.duration,"total_marks":e.total_marks,"passing_percentage":e.passing_percentage,"start_at":e.start_at.isoformat() if e.start_at else None,"end_at":e.end_at.isoformat() if e.end_at else None,"status":e.status,"questions":len(e.questions),"created_at":e.created_at.isoformat(),"security_policy":{"tab_switch_limit":e.tab_switch_limit,"enforce_fullscreen":e.enforce_fullscreen,"allow_copy_paste":e.allow_copy_paste,"show_terminated_score":e.show_terminated_score,"warning_text":e.security_warning_text,"disconnection_grace_seconds":e.disconnection_grace_seconds},"result_policy":{"publish_mode":e.result_publish_mode,"published_at":e.results_published_at.isoformat() if e.results_published_at else None,"hide_answers_until_close":e.hide_answers_until_close},"randomization":{"questions":e.randomize_questions,"options":e.randomize_options,"version_count":e.version_count}}
    if include_answers:d["assigned_student_ids"]=assigned_ids
    if detail:
        d["question_list"]=[{"id":q.id,"type":q.question_type,"text":q.question_text,"marks":q.marks,"order":q.question_order,"topic":q.topic,"difficulty":q.difficulty,"tags":q.tags or [],"options":[{"id":o.id,"text":o.option_text,**({"correct":o.is_correct} if include_answers else {})} for o in q.options],**({"correct_answer":next((o.option_text for o in q.options if o.is_correct),"")} if include_answers and q.question_type in ("TRUE_FALSE","SHORT_ANSWER","PARAGRAPH") else {})} for q in e.questions]
    return d

@api.post("/auth/<role>/login")
@limiter.limit("10 per minute")
def login(role):
    if role not in ("admin","student"):return fail("Invalid role")
    body=request.get_json() or {}; identity=(body.get("identifier") or body.get("email") or "").strip().lower()
    user=User.query.filter(func.lower(User.email)==identity).first()
    if role=="student" and not user:
        p=StudentProfile.query.filter(func.lower(StudentProfile.student_id)==identity).first(); user=p.user if p else None
    if not user or user.role!=role.upper() or not user.check_password(body.get("password", "")):
        if user:db.session.add(AuditLog(institute_id=user.institute_id,actor_id=user.id,action="LOGIN_FAILED",category="SECURITY",severity="WARNING",entity_type="USER",entity_id=user.id,details={"role":role},ip_address=client_ip(),user_agent=request.user_agent.string[:500],request_id=request.headers.get("X-Request-ID") or str(uuid.uuid4())));db.session.commit()
        return fail("Invalid credentials",401)
    if not user.email_verified:return fail("Verify your email address before signing in",403)
    if user.status=="PENDING":return fail("Your registration is waiting for admin approval",403)
    if user.status=="REJECTED":return fail("Your registration was not approved. Contact the administrator",403)
    if user.status!="ACTIVE":return fail("Your account is not active",403)
    token=create_access_token(str(user.id),additional_claims={"role":user.role,"institute_id":user.institute_id});db.session.add(AuditLog(institute_id=user.institute_id,actor_id=user.id,action="LOGIN_SUCCESS",category="SECURITY",severity="INFO",entity_type="USER",entity_id=user.id,details={"role":role},ip_address=client_ip(),user_agent=request.user_agent.string[:500],request_id=request.headers.get("X-Request-ID") or str(uuid.uuid4())));db.session.commit()
    return ok({"token":token,"user":user_json(user)})

@api.get("/auth/institutes")
def registration_institutes():return ok([institute_json(item,False) for item in Institute.query.filter_by(status="ACTIVE").order_by(Institute.name).all()])
@api.post("/auth/institutes/onboard")
@limiter.limit("3 per hour")
def onboard_institute():
    body=request.get_json() or {};required=("institute_name","slug","admin_name","email","password")
    if any(not str(body.get(field,"")).strip() for field in required):return fail("All onboarding fields are required",422)
    slug=re.sub(r"[^a-z0-9-]","",body["slug"].strip().lower().replace(" ","-"))
    if len(slug)<3:return fail("Institute code must contain at least 3 letters or numbers",422)
    email=body["email"].strip().lower();password=str(body["password"])
    if Institute.query.filter_by(slug=slug).first():return fail("This institute code is already used",409)
    if User.query.filter(func.lower(User.email)==email).first():return fail("This email is already registered",409)
    if len(password)<8:return fail("Password must contain at least 8 characters",422)
    verification_required=bool(current_app.config.get("MAIL_SERVER"));institute=Institute(name=body["institute_name"].strip(),slug=slug,primary_color="#4169e1",subscription_ends_at=now()+timedelta(days=30));db.session.add(institute);db.session.flush()
    admin=User(name=body["admin_name"].strip(),email=email,role="ADMIN",status="ACTIVE",email_verified=not verification_required,institute_id=institute.id);admin.set_password(password);db.session.add(admin);db.session.flush()
    for name in ("Intern Student","Regular Course Student","Workshop Student","Normal Student"):db.session.add(StudentCategory(name=name,institute_id=institute.id))
    raw=create_account_token(admin,"VERIFY_EMAIL",24);db.session.commit();delivered=send_account_email(admin,"Verify your Axiom Exams email",f"Verify your email address: {current_app.config['SITE_URL']}/verify-email/{raw}") if verification_required else False
    return ok({"institute":institute_json(institute),"email_verification_required":verification_required,"email_delivery_sent":delivered},"Institute created",201)
@api.get("/auth/student/categories")
def registration_categories():
    try:institute_id=int(request.args.get("institute_id",0))
    except ValueError:return fail("Select a valid institute",422)
    return ok([{"id":category.id,"name":category.name} for category in StudentCategory.query.filter_by(status="ACTIVE",institute_id=institute_id).order_by(StudentCategory.name).all()])

@api.post("/auth/student/register")
@limiter.limit("5 per minute")
def register_student():
    body=request.get_json() or {}
    required=("name","email","password","student_id","course","student_category_id","institute_id")
    if any(not str(body.get(field,"")).strip() for field in required):return fail("All registration fields are required",422)
    email=body["email"].strip().lower();student_id=body["student_id"].strip()
    if User.query.filter(func.lower(User.email)==email).first():return fail("This email is already registered",409)
    if StudentProfile.query.filter(func.lower(StudentProfile.student_id)==student_id.lower()).first():return fail("This student ID is already registered",409)
    try:institute_id=int(body["institute_id"])
    except (TypeError,ValueError):return fail("Select a valid institute",422)
    institute=Institute.query.filter_by(id=institute_id,status="ACTIVE").first();category=StudentCategory.query.filter_by(id=body["student_category_id"],institute_id=institute_id,status="ACTIVE").first()
    if not institute:return fail("Select a valid institute",422)
    if not institute_available(institute):return fail("This institute subscription is not active",403)
    if User.query.filter_by(institute_id=institute_id,role="STUDENT").count()>=institute.student_limit:return fail("This institute has reached its student limit",409)
    if not category:return fail("Select a valid student type",422)
    password=str(body["password"])
    if len(password)<8:return fail("Password must contain at least 8 characters",422)
    verification_required=bool(current_app.config.get("MAIL_SERVER"))
    user=User(name=body["name"].strip(),email=email,role="STUDENT",status="PENDING",email_verified=not verification_required,institute_id=institute_id);user.set_password(password)
    db.session.add(user);db.session.flush();db.session.add(StudentProfile(user_id=user.id,student_id=student_id,course=body["course"].strip(),student_category_id=category.id,enrollment_date=now().date()))
    raw=create_account_token(user,"VERIFY_EMAIL",24);db.session.commit()
    delivered=send_account_email(user,"Verify your Axiom Exams email",f"Verify your email address: {current_app.config['SITE_URL']}/verify-email/{raw}") if verification_required else False
    return ok({"status":"PENDING","email_verification_required":verification_required,"email_delivery_sent":delivered},"Registration submitted. Verify your email and wait for administrator approval." if verification_required else "Registration submitted. You can log in after admin approval.",201)

def token_digest(raw):return hashlib.sha256(raw.encode()).hexdigest()
def create_account_token(user,purpose,hours):
    AccountToken.query.filter_by(user_id=user.id,purpose=purpose,used_at=None).update({"used_at":now()})
    raw=secrets.token_urlsafe(32);db.session.add(AccountToken(user_id=user.id,purpose=purpose,token_hash=token_digest(raw),expires_at=now()+timedelta(hours=hours)));return raw
def valid_account_token(raw,purpose):
    token=AccountToken.query.filter_by(token_hash=token_digest(raw),purpose=purpose,used_at=None).first()
    if not token:return None
    expires=token.expires_at.replace(tzinfo=timezone.utc) if token.expires_at.tzinfo is None else token.expires_at
    return token if expires>now() else None
def send_account_email(user,subject,body):
    try:mail.send(Message(subject=subject,recipients=[user.email],body=body,sender=(current_app.config["MAIL_FROM_NAME"],current_app.config["MAIL_FROM"])));return True
    except Exception:current_app.logger.exception("Account email delivery failed");return False

@api.post("/auth/resend-verification")
@limiter.limit("3 per minute")
def resend_verification():
    email=(request.get_json() or {}).get("email","").strip().lower();user=User.query.filter(func.lower(User.email)==email,User.role=="STUDENT",User.email_verified.is_(False)).first()
    if user and current_app.config.get("MAIL_SERVER"):
        raw=create_account_token(user,"VERIFY_EMAIL",24);db.session.commit();send_account_email(user,"Verify your Axiom Exams email",f"Verify your email address: {current_app.config['SITE_URL']}/verify-email/{raw}")
    return ok(message="If the account needs verification, a new link has been sent.")

@api.post("/auth/forgot-password")
@limiter.limit("5 per minute")
def forgot_password():
    email=(request.get_json() or {}).get("email","").strip().lower();user=User.query.filter(func.lower(User.email)==email).first()
    if user and current_app.config.get("MAIL_SERVER"):
        raw=create_account_token(user,"RESET_PASSWORD",1);db.session.commit();send_account_email(user,"Reset your Axiom Exams password",f"Reset your password: {current_app.config['SITE_URL']}/reset-password/{raw}\n\nThis link expires in one hour.")
    return ok(message="If that email is registered, a password reset link has been sent.")

@api.post("/auth/reset-password/<token>")
@limiter.limit("10 per minute")
def reset_password(token):
    item=valid_account_token(token,"RESET_PASSWORD");password=str((request.get_json() or {}).get("password", ""))
    if not item:return fail("This password reset link is invalid or expired",400)
    if len(password)<8:return fail("Password must contain at least 8 characters",422)
    item.user.set_password(password);item.used_at=now();db.session.add(AuditLog(institute_id=item.user.institute_id,actor_id=item.user.id,action="PASSWORD_RESET",category="SECURITY",severity="INFO",entity_type="USER",entity_id=item.user.id,ip_address=client_ip(),user_agent=request.user_agent.string[:500],request_id=request.headers.get("X-Request-ID") or str(uuid.uuid4())));db.session.commit();return ok(message="Password reset successfully")

@api.post("/auth/verify-email/<token>")
@limiter.limit("10 per minute")
def verify_email(token):
    item=valid_account_token(token,"VERIFY_EMAIL")
    if not item:return fail("This verification link is invalid or expired",400)
    item.user.email_verified=True;item.used_at=now();db.session.add(AuditLog(institute_id=item.user.institute_id,actor_id=item.user.id,action="EMAIL_VERIFIED",category="SECURITY",severity="INFO",entity_type="USER",entity_id=item.user.id,ip_address=client_ip(),user_agent=request.user_agent.string[:500],request_id=request.headers.get("X-Request-ID") or str(uuid.uuid4())));db.session.commit();return ok(message="Email verified successfully")
@api.get("/auth/me")
@jwt_required()
def me(): return ok(user_json(db.session.get(User,uid())))
@api.post("/auth/logout")
@jwt_required()
def logout(): return ok(message="Logged out")
def user_json(u):
    d={"id":u.id,"name":u.name,"email":u.email,"role":u.role,"status":u.status,"email_verified":u.email_verified,"institute":institute_json(u.institute)}
    if u.profile:d.update({"student_id":u.profile.student_id,"course":u.profile.course,"category":u.profile.category.name,"category_id":u.profile.student_category_id,"enrollment_date":u.profile.enrollment_date.isoformat()})
    return d
def institute_logo_url(item):
    if not item.logo_url:return None
    return item.logo_url if item.logo_url.startswith(("http://","https://")) else f"{request.host_url.rstrip('/')}/api/uploads/{item.logo_url}"
def institute_available(item):
    if item.subscription_status not in ("ACTIVE","TRIAL"):return False
    if not item.subscription_ends_at:return True
    end=item.subscription_ends_at.replace(tzinfo=timezone.utc) if item.subscription_ends_at.tzinfo is None else item.subscription_ends_at
    return end>now()
def institute_json(item,include_usage=True):
    data={"id":item.id,"name":item.name,"slug":item.slug,"status":item.status,"logo_url":institute_logo_url(item),"primary_color":item.primary_color,"subscription_plan":item.subscription_plan,"subscription_status":item.subscription_status,"student_limit":item.student_limit,"exam_limit":item.exam_limit,"subscription_ends_at":item.subscription_ends_at.isoformat() if item.subscription_ends_at else None}
    if include_usage:data["usage"]={"students":User.query.filter_by(institute_id=item.id,role="STUDENT").count(),"exams":Exam.query.filter(Exam.institute_id==item.id,Exam.archived_at.is_(None)).count(),"administrators":User.query.filter_by(institute_id=item.id,role="ADMIN",status="ACTIVE").count()}
    return data

@api.get("/uploads/<path:filename>")
def uploaded_file(filename):return send_from_directory(current_app.config["UPLOAD_FOLDER"],secure_filename(filename),max_age=86400)

@api.get("/student/dashboard")
@role_required("STUDENT")
def student_dashboard():
    attempts=ExamAttempt.query.filter_by(student_id=uid()).all(); done=[a for a in attempts if a.status=="COMPLETED"]
    available=eligible_exams(); scores=[a.percentage for a in done]
    return ok({"profile":user_json(db.session.get(User,uid())),"stats":{"total_exams":len(available)+len(done),"attempted":len(attempts),"completed":len(done),"pending":len(available),"average_score":round(sum(scores)/len(scores),1) if scores else 0,"highest_score":max(scores,default=0),"pass_rate":round(100*sum(a.result=="PASS" for a in done)/len(done),1) if done else 0},"upcoming":[exam_json(e) for e in available[:4]],"recent":[attempt_json(a) for a in done[-5:]][::-1]})
def eligible_exams():
    p=StudentProfile.query.filter_by(user_id=uid()).first()
    if not p:return []
    q=Exam.query.options(joinedload(Exam.category),selectinload(Exam.assignments),selectinload(Exam.questions)).filter(Exam.institute_id==tenant_id(),Exam.status=="PUBLISHED",Exam.archived_at.is_(None),or_(Exam.student_category_id==p.student_category_id,Exam.assignments.any(ExamAssignment.student_id==uid()))).order_by(Exam.start_at)
    latest_by_exam={}
    for attempt in ExamAttempt.query.filter_by(student_id=uid()).order_by(ExamAttempt.exam_id,ExamAttempt.attempt_number.desc()).all():latest_by_exam.setdefault(attempt.exam_id,attempt)
    visible=[]
    for e in q.all():
        sync(e);latest=latest_by_exam.get(e.id)
        if e.status=="PUBLISHED" and not (latest and latest.status in ("COMPLETED","TERMINATED") and not latest.retake_allowed):visible.append(e)
    return visible
def exam_assigned_to_student(exam,student_id,profile=None):
    profile=profile or StudentProfile.query.filter_by(user_id=student_id).first()
    return bool(profile and (exam.student_category_id==profile.student_category_id or any(assignment.student_id==student_id for assignment in exam.assignments)))
@api.get("/student/exams")
@role_required("STUDENT")
def student_exams():
    items=[]
    for exam in eligible_exams():
        data=exam_json(exam)
        attempt=ExamAttempt.query.filter_by(exam_id=exam.id,student_id=uid()).order_by(ExamAttempt.attempt_number.desc()).first()
        data["attempt_state"]="RESUME" if attempt and attempt.status in ("IN_PROGRESS","PAUSED") else "NEW"
        items.append(data)
    return ok(items)
@api.get("/student/exams/code/<code>")
@role_required("STUDENT")
def student_exam_by_code(code):
    exam=next((item for item in eligible_exams() if item.access_code==code.strip().upper()),None)
    return ok(exam_json(exam)) if exam else fail("This exam link is unavailable, not published, or not assigned to your category",403)
@api.get("/student/exams/<int:eid>")
@role_required("STUDENT")
def student_exam(eid):
    e=next((x for x in eligible_exams() if x.id==eid),None)
    return ok(exam_json(e,True)) if e else fail("Exam unavailable or not eligible",403)
@api.post("/student/exams/<int:eid>/start")
@role_required("STUDENT")
def start_exam(eid):
    e=Exam.query.filter_by(id=eid,institute_id=tenant_id(),archived_at=None).first();student_id=uid();t=now()
    if not e or not exam_assigned_to_student(e,student_id):return fail("Exam unavailable or not eligible",403)
    completed=ExamAttempt.query.filter_by(exam_id=e.id,student_id=student_id,status="COMPLETED").order_by(ExamAttempt.attempt_number.desc()).first()
    if completed and not completed.retake_allowed:return ok({"completed":True,"result_id":completed.id},"You have already completed this exam")
    other_active=ExamAttempt.query.filter(ExamAttempt.student_id==student_id,ExamAttempt.exam_id!=e.id,ExamAttempt.status.in_(("IN_PROGRESS","PAUSED"))).first()
    if other_active:return fail("Finish your active exam before starting another one",409)
    body=request.get_json(silent=True) or {};session_id=str(body.get("session_id", ""))[:80];a=ExamAttempt.query.filter_by(exam_id=e.id,student_id=student_id).order_by(ExamAttempt.attempt_number.desc()).first()
    if a and a.status=="COMPLETED" and not a.retake_allowed:return ok({"completed":True,"result_id":a.id},"Exam already completed")
    if a and a.status=="TERMINATED" and not a.retake_allowed:return ok({"terminated":True,"attempt_id":a.id},"Attempt is terminated")
    if a and a.status in ("COMPLETED","TERMINATED") and a.retake_allowed:
        a.retake_allowed=False;a=ExamAttempt(exam_id=e.id,student_id=student_id,attempt_number=a.attempt_number+1,total_questions=len(e.questions),total_marks=e.total_marks,session_id=session_id,initial_ip=client_ip(),user_agent=request.user_agent.string[:500]);db.session.add(a);db.session.commit()
    if not a:
        sync(e)
        if e.status!="PUBLISHED":return fail("Exam is not available",409)
        if e.start_at and e.start_at.replace(tzinfo=timezone.utc)>t:return fail("Exam has not started",409)
        a=ExamAttempt(exam_id=e.id,student_id=student_id,total_questions=len(e.questions),total_marks=e.total_marks,session_id=session_id,initial_ip=client_ip(),user_agent=request.user_agent.string[:500]);db.session.add(a);db.session.commit()
    elif a.status=="IN_PROGRESS":
        a.refresh_count+=1
        if session_id and a.session_id and session_id!=a.session_id:record_attempt_event(a,"MULTIPLE_SESSION",20,{"previous_session":a.session_id,"new_session":session_id})
        if session_id:a.session_id=session_id
        if client_ip()!=a.initial_ip:record_attempt_event(a,"IP_CHANGE",15,{"initial_ip":a.initial_ip,"current_ip":client_ip()})
        a.last_activity_at=t;db.session.commit()
    expires_at=attempt_deadline(a)
    if t>=expires_at:
        a.auto_submitted=True;submit_attempt(a)
        return ok({"completed":True,"result_id":a.id},"Exam time expired and the attempt was submitted")
    saved_answers={str(answer.question_id):answer.answer for answer in a.answers}
    return ok({"attempt_id":a.id,"started_at":a.started_at.isoformat(),"expires_at":expires_at.isoformat(),"saved_answers":saved_answers,"paused":a.status=="PAUSED","warning":a.warning_message,"exam":exam_json(e,True)})

@api.get("/student/exams/<int:eid>/active-attempt")
@role_required("STUDENT")
def active_attempt(eid):
    a=ExamAttempt.query.filter(ExamAttempt.exam_id==eid,ExamAttempt.student_id==uid(),ExamAttempt.status.in_(("IN_PROGRESS","PAUSED"))).order_by(ExamAttempt.attempt_number.desc()).first()
    if not a:return ok(None)
    return ok({"attempt_id":a.id,"status":a.status,"expires_at":attempt_deadline(a).isoformat(),"warning":a.warning_message,"warning_sent_at":a.warning_sent_at.isoformat() if a.warning_sent_at else None})
def attempt_deadline(attempt):
    deadline=attempt.started_at.replace(tzinfo=timezone.utc)+timedelta(minutes=attempt.exam.duration)
    if attempt.exam.end_at:deadline=min(deadline,attempt.exam.end_at.replace(tzinfo=timezone.utc))
    paused=attempt.paused_seconds or 0
    if attempt.paused_at:
        paused_at=attempt.paused_at.replace(tzinfo=timezone.utc) if attempt.paused_at.tzinfo is None else attempt.paused_at;paused+=max(0,int((now()-paused_at).total_seconds()))
    return deadline+timedelta(minutes=attempt.extra_time_minutes or 0,seconds=paused)
def record_attempt_event(attempt,event_type,risk_points=0,details=None,severity=None):
    severity=severity or ("HIGH" if risk_points>=15 else "WARNING" if risk_points else "INFO");attempt.risk_score=(attempt.risk_score or 0)+risk_points;attempt.last_activity_at=now();db.session.add(AttemptEvent(attempt_id=attempt.id,event_type=event_type,severity=severity,risk_points=risk_points,details=details or {},ip_address=client_ip(),user_agent=request.user_agent.string[:500]))
@api.post("/student/attempts/<int:aid>/answers")
@role_required("STUDENT")
def save_answer(aid):
    a=ExamAttempt.query.filter_by(id=aid,student_id=uid()).first()
    if not a or a.status!="IN_PROGRESS":return fail("Attempt is not active",409)
    if now()>=attempt_deadline(a):
        a.auto_submitted=True;submit_attempt(a)
        return ok({"saved":False,"completed":True,"result_id":a.id},"Exam time expired and the attempt was submitted")
    b=request.get_json() or {}; q=Question.query.filter_by(id=b.get("question_id"),exam_id=a.exam_id).first()
    if not q:return fail("Question is not part of this exam",422)
    ans=AttemptAnswer.query.filter_by(attempt_id=a.id,question_id=q.id).first() or AttemptAnswer(attempt_id=a.id,question_id=q.id)
    ans.answer=b.get("answer"); ans.answered_at=now();a.last_activity_at=now();db.session.add(ans); db.session.commit(); return ok({"saved":True,"saved_at":ans.answered_at.isoformat()})
@api.post("/student/attempts/<int:aid>/heartbeat")
@role_required("STUDENT")
def attempt_heartbeat(aid):
    a=ExamAttempt.query.filter_by(id=aid,student_id=uid()).first()
    if not a:return fail("Attempt not found",404)
    if a.status=="COMPLETED":return ok({"completed":True,"result_id":a.id,"status":a.status})
    if a.status=="TERMINATED":return ok({"completed":False,"status":a.status})
    if a.status=="IN_PROGRESS" and now()>=attempt_deadline(a):a.auto_submitted=True;submit_attempt(a);return ok({"completed":True,"result_id":a.id,"status":a.status})
    body=request.get_json(silent=True) or {};session_id=str(body.get("session_id", ""))[:80]
    if session_id and a.session_id and session_id!=a.session_id:
        record_attempt_event(a,"MULTIPLE_SESSION",20,{"blocked_session":session_id},"HIGH");db.session.commit();return fail("This exam is already active in another browser session",409)
    if session_id and not a.session_id:a.session_id=session_id
    previous=a.last_activity_at
    if previous:
        previous=previous.replace(tzinfo=timezone.utc) if previous.tzinfo is None else previous
        if (now()-previous).total_seconds()>60:record_attempt_event(a,"INACTIVITY",8,{"inactive_seconds":int((now()-previous).total_seconds())})
    if client_ip()!=a.initial_ip:record_attempt_event(a,"IP_CHANGE",15,{"initial_ip":a.initial_ip,"current_ip":client_ip()})
    presence={"attempt_id":a.id,"exam_id":a.exam_id,"student_id":a.student_id,"status":a.status,"last_seen":now().isoformat()};stored=set_json(f"exam:{a.exam_id}:student:{a.student_id}:presence",presence,current_app.config["PRESENCE_TTL_SECONDS"])
    if not stored:a.last_activity_at=now()
    db.session.commit();return ok({"status":a.status,"paused":a.status=="PAUSED","expires_at":attempt_deadline(a).isoformat(),"warning":a.warning_message,"warning_sent_at":a.warning_sent_at.isoformat() if a.warning_sent_at else None,"heartbeat_interval":current_app.config["HEARTBEAT_INTERVAL_SECONDS"]})
@api.post("/student/attempts/<int:aid>/events")
@role_required("STUDENT")
def attempt_event(aid):
    a=ExamAttempt.query.filter_by(id=aid,student_id=uid()).with_for_update().first();body=request.get_json() or {};event_type=str(body.get("event_type","")).upper()
    if not a or a.status not in ("IN_PROGRESS","PAUSED"):return fail("Attempt is not active",409)
    weights={"TAB_HIDDEN":4,"FULLSCREEN_EXIT":6,"COPY":3,"PASTE":3,"CONTEXT_MENU":2,"REFRESH":3,"PROLONGED_DISCONNECTION":10}
    if event_type not in weights:return fail("Unsupported monitoring event",422)
    if event_type=="REFRESH":a.refresh_count+=1
    if event_type=="TAB_HIDDEN":
        switch_count=AttemptEvent.query.filter_by(attempt_id=a.id,event_type="TAB_HIDDEN").count()+1
        limit=max(0,a.exam.tab_switch_limit)
        record_attempt_event(a,event_type,weights[event_type],{"switch_count":switch_count,"allowed":limit},"HIGH" if switch_count>limit else "WARNING")
        if switch_count<=limit:
            db.session.commit()
            message=a.exam.security_warning_text or f"Warning: tab switch {switch_count} of {limit}. The next violation terminates and auto-submits your exam."
            return ok({"recorded":True,"risk_score":a.risk_score,"action":"WARNING","message":message,"switch_count":switch_count,"limit":limit})
        a.auto_submitted=True;a.termination_reason=f"Automatically terminated after exceeding the allowed tab-switch limit ({limit}).";a.warning_message=a.termination_reason;a.warning_sent_at=now()
        audit("AUTO_TERMINATE_ATTEMPT",a.id,"ATTEMPT",{"reason":a.termination_reason,"tab_switches":switch_count},"SECURITY","HIGH")
        submit_attempt(a,final_status="TERMINATED",message="Exam auto-submitted")
        return ok({"action":"TERMINATED","attempt_id":a.id},"Exam auto-submitted")
    record_attempt_event(a,event_type,weights[event_type],body.get("details"));db.session.commit();return ok({"recorded":True,"risk_score":a.risk_score})
@api.post("/student/attempts/<int:aid>/submit")
@role_required("STUDENT")
def submit(aid):
    lock_key=f"submit:{aid}";locked=acquire_lock(lock_key,60)
    a=ExamAttempt.query.filter_by(id=aid,student_id=uid()).with_for_update().first()
    if not a:return fail("Attempt not found",404)
    if a.status=="COMPLETED":return ok(attempt_json(a,True),"Already submitted")
    if not locked:return fail("Submission is already being processed",409)
    try:return submit_attempt(a)
    finally:release_lock(lock_key)
def submit_attempt(a,final_status="COMPLETED",message="Exam submitted"):
    amap={x.question_id:x for x in a.answers}; correct=wrong=attempted=0; marks=0
    for q in a.exam.questions:
        x=amap.get(q.id)
        if not x:continue
        val=x.answer; chosen=set(map(str,val if isinstance(val,list) else [val])); expected={str(o.id) for o in q.options if o.is_correct}
        if q.question_type in ("SHORT_ANSWER", "PARAGRAPH"):
            expected={o.option_text.strip().casefold() for o in q.options if o.is_correct}
            chosen={str(val).strip().casefold()}
        x.is_correct=chosen==expected; x.marks_obtained=q.marks if x.is_correct else 0; attempted+=1; correct+=int(x.is_correct); wrong+=int(not x.is_correct); marks+=x.marks_obtained
    a.status=final_status; a.submitted_at=now(); a.attempted_questions=attempted;a.correct_answers=correct;a.wrong_answers=wrong;a.unanswered_questions=len(a.exam.questions)-attempted;a.obtained_marks=marks;a.percentage=round(100*marks/a.total_marks,2) if a.total_marks else 0;a.result="PASS" if a.percentage>=a.exam.passing_percentage else "FAIL";db.session.commit()
    data=attempt_json(a,True)
    if final_status=="TERMINATED":data["action"]="TERMINATED"
    return ok(data,message)
def attempt_json(a,detail=False):
    d={"id":a.id,"exam_id":a.exam_id,"exam_name":a.exam.title,"exam_category":a.exam.exam_category,"student":a.student.name,"student_id":a.student.profile.student_id if a.student.profile else None,"started_at":a.started_at.isoformat(),"submitted_at":a.submitted_at.isoformat() if a.submitted_at else None,"status":a.status,"total_questions":a.total_questions,"attempted_questions":a.attempted_questions,"correct_answers":a.correct_answers,"wrong_answers":a.wrong_answers,"unanswered_questions":a.unanswered_questions,"total_marks":a.total_marks,"obtained_marks":a.obtained_marks,"percentage":a.percentage,"result":a.result,"time_taken":int((a.submitted_at-a.started_at).total_seconds()) if a.submitted_at else None}
    if detail:d["answers"]=[{"question":x.question.question_text,"answer":x.answer,"correct_answer":[o.option_text for o in x.question.options if o.is_correct],"is_correct":x.is_correct,"marks":x.marks_obtained} for x in a.answers]
    return d

def answer_display(question,answer):
    if answer in (None,""):return "Not answered"
    if question.question_type in ("SHORT_ANSWER","PARAGRAPH"):return str(answer)
    chosen=answer if isinstance(answer,list) else [answer]
    option_map={str(option.id):option.option_text for option in question.options}
    return ", ".join(option_map.get(str(value),str(value)) for value in chosen) or "Not answered"

@api.get("/admin/exams/<int:eid>/report")
@role_required("ADMIN")
def exam_report(eid):
    exam=Exam.query.filter_by(id=eid,institute_id=tenant_id(),archived_at=None).first()
    if not exam:return fail("Exam not found",404)
    attempts=ExamAttempt.query.filter_by(exam_id=eid).order_by(ExamAttempt.started_at.desc()).all()
    rows=[]
    for attempt in attempts:
        answers={answer.question_id:answer for answer in attempt.answers}
        review=[]
        for question in exam.questions:
            answer=answers.get(question.id)
            review.append({"question_id":question.id,"question":question.question_text,"question_type":question.question_type,"total_marks":question.marks,"selected_answer":answer_display(question,answer.answer) if answer else "Not answered","correct_answer":", ".join(option.option_text for option in question.options if option.is_correct),"is_correct":answer.is_correct if answer else False,"marks_obtained":answer.marks_obtained if answer else 0})
        item=attempt_json(attempt);item["answers"]=review;rows.append(item)
    completed=[attempt for attempt in attempts if attempt.status=="COMPLETED"]
    return ok({"exam":exam_json(exam),"summary":{"attended":len(attempts),"completed":len(completed),"in_progress":len(attempts)-len(completed),"passed":sum(attempt.result=="PASS" for attempt in completed),"failed":sum(attempt.result=="FAIL" for attempt in completed),"average_mark":round(sum(attempt.obtained_marks for attempt in completed)/len(completed),2) if completed else 0},"attempts":rows})
@api.get("/student/history")
@role_required("STUDENT")
def history():
    items=[]
    for attempt in ExamAttempt.query.filter_by(student_id=uid()).order_by(ExamAttempt.started_at.desc()).all():
        if attempt.status=="TERMINATED":items.append({"id":attempt.id,"exam_id":attempt.exam_id,"exam_name":attempt.exam.title,"exam_category":attempt.exam.exam_category,"started_at":attempt.started_at.isoformat(),"submitted_at":attempt.submitted_at.isoformat() if attempt.submitted_at else None,"status":"TERMINATED","result":None,"percentage":None})
        else:items.append(attempt_json(attempt))
    return ok(items)
@api.get("/student/attempts/<int:aid>/termination")
@role_required("STUDENT")
def student_termination(aid):
    attempt=ExamAttempt.query.filter_by(id=aid,student_id=uid(),status="TERMINATED").first()
    if not attempt:return fail("Terminated attempt not found",404)
    return ok({"id":attempt.id,"exam_name":attempt.exam.title,"status":attempt.status,"submitted_at":attempt.submitted_at.isoformat() if attempt.submitted_at else None,"retake_requested":False})
@api.get("/student/results/<int:aid>")
@role_required("STUDENT")
def result(aid):
    a=ExamAttempt.query.filter_by(id=aid,student_id=uid(),status="COMPLETED").first(); return ok(attempt_json(a,True)) if a else fail("Result not found",404)

@api.get("/admin/dashboard")
@role_required("ADMIN")
def admin_dashboard():
    exams=Exam.query.filter(Exam.institute_id==tenant_id(),Exam.archived_at.is_(None)).all();query=ExamAttempt.query.join(Exam).join(User,ExamAttempt.student_id==User.id).outerjoin(StudentProfile,StudentProfile.user_id==User.id).filter(Exam.institute_id==tenant_id())
    if request.args.get("exam_id"):query=query.filter(ExamAttempt.exam_id==int(request.args["exam_id"]))
    if request.args.get("category_id"):query=query.filter(StudentProfile.student_category_id==int(request.args["category_id"]))
    if request.args.get("course"):query=query.filter(StudentProfile.course==request.args["course"])
    if request.args.get("from"):query=query.filter(ExamAttempt.started_at>=dt(request.args["from"]))
    if request.args.get("to"):query=query.filter(ExamAttempt.started_at<=dt(request.args["to"]))
    all_attempts=query.all();ats=[a for a in all_attempts if a.status=="COMPLETED"];scores=[a.percentage for a in ats];active=[a for a in all_attempts if a.status in ("IN_PROGRESS","PAUSED")];online_cutoff=now()-timedelta(seconds=45);online=[a for a in active if a.last_activity_at and (a.last_activity_at.replace(tzinfo=timezone.utc) if a.last_activity_at.tzinfo is None else a.last_activity_at)>=online_cutoff]
    by_exam=[]
    for exam in exams:
        rows=[a for a in ats if a.exam_id==exam.id];fail_rate=round(100*sum(a.result=="FAIL" for a in rows)/len(rows),1) if rows else 0
        if len(rows)>=3 and fail_rate>=70:by_exam.append({"exam_id":exam.id,"exam":exam.title,"attempts":len(rows),"failure_rate":fail_rate})
    student_rows={}
    for a in ats:student_rows.setdefault(a.student_id,[]).append(a.percentage)
    support=[{"student_id":sid,"student":db.session.get(User,sid).name,"average":round(sum(values)/len(values),1)} for sid,values in student_rows.items() if sum(values)/len(values)<50]
    comparisons={}
    for a in ats:
        profile=a.student.profile;key=(profile.course if profile else None) or "Unassigned";comparisons.setdefault(key,[]).append(a.percentage)
    institute=db.session.get(Institute,tenant_id());student_count=User.query.filter_by(role="STUDENT",institute_id=tenant_id()).count();completion=round(100*len(ats)/len(all_attempts),1) if all_attempts else 0;pass_rate=round(100*sum(a.result=="PASS" for a in ats)/len(ats),1) if ats else 0
    return ok({"stats":{"students":student_count,"exams":len(exams),"draft":sum(e.status=="DRAFT" for e in exams),"published":sum(sync(e).status=="PUBLISHED" for e in exams),"completed":sum(e.status=="COMPLETED" for e in exams),"attempts":len(ats),"passed":sum(a.result=="PASS" for a in ats),"failed":sum(a.result=="FAIL" for a in ats),"average_score":round(sum(scores)/len(scores),1) if scores else 0,"active_attempts":len(active),"students_online":len(online),"completion_rate":completion,"pass_rate":pass_rate,"abandoned":sum(a.status=="TERMINATED" for a in all_attempts)},"unusual_failure_rates":by_exam,"students_needing_support":support,"course_comparisons":[{"course":key,"attempts":len(values),"average":round(sum(values)/len(values),1)} for key,values in comparisons.items()],"warnings":{"student_usage":round(100*student_count/institute.student_limit,1) if institute.student_limit else 100,"exam_usage":round(100*len(exams)/institute.exam_limit,1) if institute.exam_limit else 100,"subscription_status":institute.subscription_status,"subscription_ends_at":institute.subscription_ends_at.isoformat() if institute.subscription_ends_at else None},"health":{"api":"HEALTHY","database":"HEALTHY","checked_at":now().isoformat()},"filters":{"exams":[{"id":e.id,"title":e.title} for e in exams],"courses":sorted({u.profile.course for u in User.query.filter_by(role="STUDENT",institute_id=tenant_id()).all() if u.profile and u.profile.course}),"categories":[{"id":c.id,"name":c.name} for c in StudentCategory.query.filter_by(institute_id=tenant_id()).all()]},"recent_attempts":[attempt_json(a) for a in ats[-6:]][::-1]})
@api.route("/admin/institute",methods=["GET","PUT"])
@role_required("ADMIN")
def institute_settings():
    institute=db.session.get(Institute,tenant_id())
    if request.method=="GET":return ok(institute_json(institute))
    body=request.get_json() or {};name=str(body.get("name","")).strip();color=str(body.get("primary_color","")).strip()
    if not name:return fail("Institute name is required",422)
    if not re.fullmatch(r"#[0-9a-fA-F]{6}",color):return fail("Choose a valid brand colour",422)
    institute.name=name[:160];institute.primary_color=color.lower();db.session.commit();return ok(institute_json(institute),"Institute settings updated")
@api.post("/admin/institute/logo")
@role_required("ADMIN")
def upload_institute_logo():
    upload=request.files.get("logo")
    if not upload or not upload.filename:return fail("Choose a logo file",422)
    header=upload.stream.read(12);upload.stream.seek(0)
    kinds=[("png",header.startswith(b"\x89PNG\r\n\x1a\n")),("jpg",header.startswith(b"\xff\xd8\xff")),("webp",header.startswith(b"RIFF") and header[8:12]==b"WEBP")];extension=next((name for name,valid in kinds if valid),None)
    if not extension:return fail("Logo must be a PNG, JPEG, or WebP image",422)
    folder=current_app.config["UPLOAD_FOLDER"];os.makedirs(folder,exist_ok=True);filename=f"institute-{tenant_id()}-{uuid.uuid4().hex}.{extension}";upload.save(os.path.join(folder,filename))
    institute=db.session.get(Institute,tenant_id());institute.logo_url=filename;db.session.commit();return ok(institute_json(institute),"Logo updated")
@api.route("/admin/institute/administrators",methods=["GET","POST"])
@role_required("ADMIN")
def institute_administrators():
    if request.method=="GET":return ok([user_json(user) for user in User.query.filter_by(institute_id=tenant_id(),role="ADMIN").order_by(User.created_at).all()])
    body=request.get_json() or {};name=str(body.get("name","")).strip();email=str(body.get("email","")).strip().lower();password=str(body.get("password", ""))
    if not name or not email:return fail("Administrator name and email are required",422)
    if len(password)<8:return fail("Password must contain at least 8 characters",422)
    if User.query.filter(func.lower(User.email)==email).first():return fail("This email is already registered",409)
    user=User(name=name,email=email,role="ADMIN",status="ACTIVE",email_verified=True,institute_id=tenant_id());user.set_password(password);db.session.add(user);db.session.commit();return ok(user_json(user),"Administrator added",201)
@api.put("/admin/institute/administrators/<int:admin_id>/<action>")
@role_required("ADMIN")
def update_institute_administrator(admin_id,action):
    user=User.query.filter_by(id=admin_id,institute_id=tenant_id(),role="ADMIN").first()
    if not user:return fail("Administrator not found",404)
    if action not in ("suspend","reactivate"):return fail("Invalid action",404)
    if action=="suspend":
        if user.id==uid():return fail("You cannot suspend your own account",409)
        if User.query.filter_by(institute_id=tenant_id(),role="ADMIN",status="ACTIVE").count()<=1:return fail("At least one active administrator is required",409)
        user.status="SUSPENDED"
    else:user.status="ACTIVE"
    db.session.commit();return ok(user_json(user),f"Administrator {action}d")
@api.get("/admin/categories")
@role_required("ADMIN")
def categories(): return ok([{"id":c.id,"name":c.name} for c in StudentCategory.query.filter_by(status="ACTIVE",institute_id=tenant_id()).all()])

def course_json(item):return {"id":item.id,"name":item.name,"code":item.code,"description":item.description,"status":item.status,"students":StudentProfile.query.filter_by(course_id=item.id).count()}
def batch_json(item):
    attempts=ExamAttempt.query.join(User,ExamAttempt.student_id==User.id).join(StudentProfile,StudentProfile.user_id==User.id).filter(StudentProfile.batch_id==item.id,ExamAttempt.status=="COMPLETED").all()
    return {"id":item.id,"name":item.name,"course_id":item.course_id,"course":item.course.name,"trainer_name":item.trainer_name,"start_date":item.start_date.isoformat() if item.start_date else None,"end_date":item.end_date.isoformat() if item.end_date else None,"status":item.status,"students":StudentProfile.query.filter_by(batch_id=item.id).count(),"performance":{"attempts":len(attempts),"average":round(sum(a.percentage for a in attempts)/len(attempts),1) if attempts else 0,"pass_rate":round(100*sum(a.result=="PASS" for a in attempts)/len(attempts),1) if attempts else 0}}

@api.route("/admin/courses",methods=["GET","POST"])
@role_required("ADMIN")
def courses():
    if request.method=="GET":return ok([course_json(item) for item in Course.query.filter_by(institute_id=tenant_id()).order_by(Course.name).all()])
    body=request.get_json() or {};name=str(body.get("name","")).strip()
    if not name:return fail("Course name is required",422)
    if Course.query.filter(func.lower(Course.name)==name.casefold(),Course.institute_id==tenant_id()).first():return fail("Course already exists",409)
    item=Course(institute_id=tenant_id(),name=name,code=str(body.get("code","")).strip() or None,description=str(body.get("description","")).strip() or None);db.session.add(item);db.session.flush();audit("CREATE_COURSE",item.id,"COURSE",{"after":course_json(item)});db.session.commit();return ok(course_json(item),"Course created",201)

@api.route("/admin/batches",methods=["GET","POST"])
@role_required("ADMIN")
def batches():
    if request.method=="GET":return ok([batch_json(item) for item in Batch.query.filter_by(institute_id=tenant_id()).order_by(Batch.created_at.desc()).all()])
    body=request.get_json() or {};course=Course.query.filter_by(id=body.get("course_id"),institute_id=tenant_id()).first();name=str(body.get("name","")).strip()
    if not course or not name:return fail("Batch name and a valid course are required",422)
    if Batch.query.filter(func.lower(Batch.name)==name.casefold(),Batch.institute_id==tenant_id()).first():return fail("Batch already exists",409)
    try:start=datetime.fromisoformat(body["start_date"]).date() if body.get("start_date") else None;end=datetime.fromisoformat(body["end_date"]).date() if body.get("end_date") else None
    except ValueError:return fail("Use valid batch dates",422)
    item=Batch(institute_id=tenant_id(),course_id=course.id,name=name,start_date=start,end_date=end,trainer_name=str(body.get("trainer_name","")).strip() or None);db.session.add(item);db.session.flush();audit("CREATE_BATCH",item.id,"BATCH",{"after":batch_json(item)});db.session.commit();return ok(batch_json(item),"Batch created",201)

@api.put("/admin/students/<int:student_id>/batch")
@role_required("ADMIN")
def transfer_student(student_id):
    student=User.query.filter_by(id=student_id,role="STUDENT",institute_id=tenant_id()).first();body=request.get_json() or {};batch=Batch.query.filter_by(id=body.get("batch_id"),institute_id=tenant_id()).first()
    if not student or not batch:return fail("Student or batch not found",404)
    before={"course_id":student.profile.course_id,"batch_id":student.profile.batch_id};student.profile.course_id=batch.course_id;student.profile.batch_id=batch.id;student.profile.course=batch.course.name;audit("TRANSFER_STUDENT",student.id,"USER",{"before":before,"after":{"course_id":batch.course_id,"batch_id":batch.id},"reason":body.get("reason")});db.session.commit();return ok(user_json(student),"Student transferred")
@api.get("/admin/students")
@role_required("ADMIN")
def students(): return ok([user_json(u) for u in User.query.filter_by(role="STUDENT",institute_id=tenant_id()).order_by(User.created_at.desc()).all()])

@api.put("/admin/students/<int:student_id>/<action>")
@role_required("ADMIN")
def verify_student(student_id,action):
    student=User.query.filter_by(id=student_id,role="STUDENT",institute_id=tenant_id()).first()
    if not student:return fail("Student not found",404)
    statuses={"approve":"ACTIVE","reject":"REJECTED"}
    if action not in statuses:return fail("Invalid verification action",404)
    if action=="approve" and not student.email_verified:return fail("The student must verify their email before approval",409)
    student.status=statuses[action];audit(f"{action.upper()}_STUDENT",student.id,"USER",{"after":{"status":student.status}});db.session.commit()
    return ok(user_json(student),f"Student {action}d successfully")

EXAM_IMPORT_HEADERS=["Exam Title","Description","Exam Category","Duration (Minutes)","Passing Percentage","Audience Category","Start At","End At","Question Type","Question","Marks","Option A","Option B","Option C","Option D","Option E","Option F","Correct Answer"]

@api.get("/admin/exams/import-template")
@role_required("ADMIN")
def exam_import_template():
    wb=Workbook();ws=wb.active;ws.title="Exam Import";ws.append(EXAM_IMPORT_HEADERS)
    examples=[
        ["Python Basics","Introductory Python assessment","Technical",30,50,"Regular Course Student","2026-09-01 09:00","2026-09-07 18:00","SINGLE","Which keyword defines a function?",2,"def","func","method","function","","","A"],
        ["Python Basics","Introductory Python assessment","Technical",30,50,"Regular Course Student","2026-09-01 09:00","2026-09-07 18:00","MULTIPLE","Select immutable types.",3,"tuple","list","string","dictionary","","","A,C"],
        ["Python Basics","Introductory Python assessment","Technical",30,50,"Regular Course Student","2026-09-01 09:00","2026-09-07 18:00","TRUE_FALSE","Python is case-sensitive.",1,"","","","","","","True"],
        ["Python Basics","Introductory Python assessment","Technical",30,50,"Regular Course Student","2026-09-01 09:00","2026-09-07 18:00","SHORT_ANSWER","What keyword exits a loop?",2,"","","","","","","break"]]
    for row in examples:ws.append(row)
    fill=PatternFill("solid",fgColor="17324D")
    for cell in ws[1]:cell.font=Font(color="FFFFFF",bold=True);cell.fill=fill;cell.alignment=Alignment(wrap_text=True)
    for i,width in enumerate([24,32,20,20,20,22,20,20,18,45,10,22,22,22,22,22,22,28],1):ws.column_dimensions[ws.cell(1,i).column_letter].width=width
    ws.freeze_panes="A2";ws.auto_filter.ref="A1:R5"
    types=DataValidation(type="list",formula1='"SINGLE,MULTIPLE,TRUE_FALSE,SHORT_ANSWER,PARAGRAPH"',allow_blank=False);ws.add_data_validation(types);types.add("I2:I1000")
    output=BytesIO();wb.save(output);output.seek(0);audit("EXPORT_EXAM_TEMPLATE",None,"EXAM",category="TRANSACTION");db.session.commit()
    return send_file(output,as_attachment=True,download_name="exam_import_template.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

def import_text(value):return str(value).strip() if value is not None else ""
def import_date(value,row,label):
    if value in (None,""):return None
    if isinstance(value,datetime):return value
    try:return datetime.fromisoformat(import_text(value).replace("Z","+00:00"))
    except ValueError:raise ValueError(f"Row {row}: {label} must be YYYY-MM-DD HH:MM")
def imported_question(values,row):
    qtype=import_text(values[8]).upper();question=import_text(values[9]);correct=import_text(values[17])
    if qtype not in {"SINGLE","MULTIPLE","TRUE_FALSE","SHORT_ANSWER","PARAGRAPH"}:raise ValueError(f"Row {row}: invalid Question Type")
    if not question:raise ValueError(f"Row {row}: Question is required")
    try:marks=float(values[10])
    except (TypeError,ValueError):raise ValueError(f"Row {row}: Marks must be a number")
    if marks<=0:raise ValueError(f"Row {row}: Marks must be greater than zero")
    item={"type":qtype,"text":question,"marks":marks}
    if qtype in ("SINGLE","MULTIPLE"):
        option_texts=[import_text(x) for x in values[11:17]];options=[x for x in option_texts if x]
        if len(options)<2:raise ValueError(f"Row {row}: at least Option A and Option B are required")
        indexes=set()
        for answer in [x.strip() for x in correct.split(",") if x.strip()]:
            if len(answer)==1 and answer.upper() in "ABCDEF":indexes.add(ord(answer.upper())-65)
            elif answer in option_texts:indexes.add(option_texts.index(answer))
            else:raise ValueError(f"Row {row}: Correct Answer must use option letters (for example A or A,C) or exact option text")
        if not indexes:raise ValueError(f"Row {row}: Correct Answer is required")
        if qtype=="SINGLE" and len(indexes)!=1:raise ValueError(f"Row {row}: SINGLE requires exactly one correct answer")
        item["options"]=[{"text":text,"correct":index in indexes} for index,text in enumerate(option_texts) if text]
    else:
        if not correct:raise ValueError(f"Row {row}: Correct Answer is required")
        if qtype=="TRUE_FALSE" and correct.casefold() not in ("true","false"):raise ValueError(f"Row {row}: TRUE_FALSE answer must be True or False")
        item["correct_answer"]=correct
    return item

@api.post("/admin/exams/import")
@role_required("ADMIN")
def import_exams():
    upload=request.files.get("file")
    if not upload or not upload.filename:return fail("Choose an Excel file to upload",422)
    if not upload.filename.lower().endswith(".xlsx"):return fail("Only .xlsx files are supported",422)
    try:
        ws=load_workbook(upload,read_only=True,data_only=True).active
        if [import_text(x.value) for x in ws[1]][:len(EXAM_IMPORT_HEADERS)]!=EXAM_IMPORT_HEADERS:raise ValueError("The spreadsheet columns do not match the downloaded template")
        groups={}
        for row_number,row in enumerate(ws.iter_rows(min_row=2,max_col=len(EXAM_IMPORT_HEADERS),values_only=True),2):
            if not any(value not in (None,"") for value in row):continue
            title=import_text(row[0])
            if not title:raise ValueError(f"Row {row_number}: Exam Title is required")
            if title not in groups:
                try:duration=int(row[3]);passing=float(row[4])
                except (TypeError,ValueError):raise ValueError(f"Row {row_number}: Duration and Passing Percentage must be numbers")
                if duration<=0 or not 0<=passing<=100:raise ValueError(f"Row {row_number}: use a positive duration and passing percentage from 0 to 100")
                audience=import_text(row[5]);category=StudentCategory.query.filter(func.lower(StudentCategory.name)==audience.casefold(),StudentCategory.status=="ACTIVE",StudentCategory.institute_id==tenant_id()).first() if audience else None
                if not category:raise ValueError(f"Row {row_number}: Audience Category '{audience}' was not found")
                groups[title]={"description":import_text(row[1]),"exam_category":import_text(row[2]),"duration":duration,"passing":passing,"category_id":category.id,"start":import_date(row[6],row_number,"Start At"),"end":import_date(row[7],row_number,"End At"),"questions":[]}
                if not groups[title]["exam_category"]:raise ValueError(f"Row {row_number}: Exam Category is required")
            groups[title]["questions"].append(imported_question(row,row_number))
        if not groups:raise ValueError("The spreadsheet has no question rows")
        institute=db.session.get(Institute,tenant_id());current_exams=Exam.query.filter(Exam.institute_id==tenant_id(),Exam.archived_at.is_(None)).count()
        if not institute_available(institute):raise ValueError("Your institute subscription is not active")
        if current_exams+len(groups)>institute.exam_limit:raise ValueError(f"Import exceeds your plan limit of {institute.exam_limit} exams")
        created=[]
        for title,data in groups.items():
            e=Exam(title=title,description=data["description"],exam_category=data["exam_category"],student_category_id=data["category_id"],institute_id=tenant_id(),duration=data["duration"],total_marks=0,passing_percentage=data["passing"],start_at=data["start"],end_at=data["end"],created_by=uid());db.session.add(e);db.session.flush();add_questions(e,data["questions"]);e.total_marks=sum(q.marks for q in e.questions);audit("IMPORT_EXAM",e.id);created.append(e)
        db.session.commit();return ok([exam_json(e,True) for e in created],f"Created {len(created)} draft exam(s) with {sum(len(e.questions) for e in created)} questions",201)
    except (ValueError,TypeError,KeyError) as error:db.session.rollback();return fail(str(error),422)
    except Exception:db.session.rollback();return fail("The Excel file could not be read. Download a fresh template and try again.",422)

@api.route("/admin/exams",methods=["GET","POST"])
@role_required("ADMIN")
def admin_exams():
    if request.method=="GET":return ok([exam_json(e) for e in Exam.query.filter(Exam.institute_id==tenant_id(),Exam.archived_at.is_(None)).order_by(Exam.created_at.desc()).all()])
    institute=db.session.get(Institute,tenant_id())
    if not institute_available(institute):return fail("Your institute subscription is not active",403)
    if Exam.query.filter(Exam.institute_id==tenant_id(),Exam.archived_at.is_(None)).count()>=institute.exam_limit:return fail(f"Your plan allows up to {institute.exam_limit} exams",409)
    b=request.get_json() or {}; required=("title","exam_category","duration","total_marks","passing_percentage")
    if any(b.get(k) in (None,"") for k in required):return fail("Missing required fields",422)
    try:category_id,student_ids=exam_audience(b)
    except ValueError as error:return fail(str(error),422)
    e=Exam(title=b["title"],description=b.get("description"),exam_category=b["exam_category"],student_category_id=category_id,institute_id=tenant_id(),duration=int(b["duration"]),total_marks=float(b["total_marks"]),passing_percentage=float(b["passing_percentage"]),start_at=dt(b.get("start_at")),end_at=dt(b.get("end_at")),created_by=uid())
    try:apply_exam_policies(e,b)
    except (ValueError,TypeError) as error:return fail(str(error),422)
    db.session.add(e);db.session.flush();set_exam_assignments(e,student_ids);add_questions(e,b.get("questions",[]));e.total_marks=sum(q.marks for q in e.questions);audit("CREATE_EXAM",e.id);db.session.commit();return ok(exam_json(e,True),"Draft created",201)
@api.route("/admin/exams/<int:eid>",methods=["GET","PUT","DELETE"])
@role_required("ADMIN")
def admin_exam(eid):
    e=Exam.query.filter_by(id=eid,institute_id=tenant_id(),archived_at=None).first()
    if not e:return fail("Exam not found",404)
    if request.method=="GET":return ok(exam_json(e,True,True))
    if request.method=="DELETE":e.archived_at=now();audit("ARCHIVE_EXAM",e.id);db.session.commit();return ok(message="Exam archived")
    b=request.get_json() or {}
    for k in ("title","description","exam_category","duration","total_marks","passing_percentage"):
        if k in b:setattr(e,k,b[k] or None)
    if any(k in b for k in ("audience_mode","student_category_id","assigned_student_ids")):
        try:e.student_category_id,student_ids=exam_audience(b);set_exam_assignments(e,student_ids)
        except ValueError as error:return fail(str(error),422)
    if "start_at" in b:e.start_at=dt(b["start_at"])
    if "end_at" in b:e.end_at=dt(b["end_at"])
    if "questions" in b:e.questions.clear();db.session.flush();add_questions(e,b["questions"]);e.total_marks=sum(q.marks for q in e.questions)
    try:apply_exam_policies(e,b)
    except (ValueError,TypeError) as error:return fail(str(error),422)
    audit("UPDATE_EXAM",e.id);db.session.commit();return ok(exam_json(e,True),"Exam updated")
def exam_audience(body):
    mode=body.get("audience_mode","category")
    if mode=="category":
        try:category_id=int(body.get("student_category_id"))
        except (TypeError,ValueError):raise ValueError("Select a student type for this exam")
        if not StudentCategory.query.filter_by(id=category_id,status="ACTIVE",institute_id=tenant_id()).first():raise ValueError("Select a valid student type")
        return category_id,[]
    if mode!="students":raise ValueError("Select a valid audience type")
    try:student_ids=list(dict.fromkeys(int(value) for value in body.get("assigned_student_ids",[])))
    except (TypeError,ValueError):raise ValueError("Select valid students")
    valid={row[0] for row in db.session.query(User.id).filter(User.id.in_(student_ids),User.role=="STUDENT",User.status=="ACTIVE",User.institute_id==tenant_id()).all()} if student_ids else set()
    if not student_ids or valid!=set(student_ids):raise ValueError("Select at least one active student")
    return None,student_ids
def set_exam_assignments(exam,student_ids):
    exam.assignments.clear();db.session.flush()
    for student_id in student_ids:exam.assignments.append(ExamAssignment(student_id=student_id,assigned_by=uid()))
def apply_exam_policies(exam,body):
    security=body.get("security_policy") or {};result=body.get("result_policy") or {};randomization=body.get("randomization") or {}
    if "tab_switch_limit" in security:exam.tab_switch_limit=max(0,min(20,int(security["tab_switch_limit"])))
    if "disconnection_grace_seconds" in security:exam.disconnection_grace_seconds=max(0,min(3600,int(security["disconnection_grace_seconds"])))
    for key,attr in (("enforce_fullscreen","enforce_fullscreen"),("allow_copy_paste","allow_copy_paste"),("show_terminated_score","show_terminated_score")):
        if key in security:setattr(exam,attr,bool(security[key]))
    if "warning_text" in security:exam.security_warning_text=str(security["warning_text"] or "")[:500] or None
    if "publish_mode" in result:
        mode=str(result["publish_mode"]).upper()
        if mode not in ("IMMEDIATE","MANUAL","AFTER_CLOSE"):raise ValueError("Invalid result publishing mode")
        exam.result_publish_mode=mode
    if "hide_answers_until_close" in result:exam.hide_answers_until_close=bool(result["hide_answers_until_close"])
    if "questions" in randomization:exam.randomize_questions=bool(randomization["questions"])
    if "options" in randomization:exam.randomize_options=bool(randomization["options"])
    if "version_count" in randomization:exam.version_count=max(1,min(20,int(randomization["version_count"])))
def add_questions(e,items):
    allowed_types={"SINGLE","MULTIPLE","TRUE_FALSE","SHORT_ANSWER","PARAGRAPH"}
    for i,item in enumerate(items):
        question_type=item.get("type","SINGLE").upper()
        if question_type not in allowed_types: raise ValueError(f"Unsupported question type: {question_type}")
        q=Question(exam=e,question_type=question_type,question_text=item["text"].strip(),marks=float(item.get("marks",1)),question_order=int(item.get("order",i+1)),topic=str(item.get("topic") or "").strip() or None,difficulty=str(item.get("difficulty") or "MEDIUM").upper(),tags=item.get("tags") or []);db.session.add(q);db.session.flush()
        options=item.get("options",[])
        if question_type=="TRUE_FALSE":
            correct=str(item.get("correct_answer", "true")).lower()=="true"
            options=[{"text":"True","correct":correct},{"text":"False","correct":not correct}]
        elif question_type in ("SHORT_ANSWER","PARAGRAPH"):
            options=[{"text":item.get("correct_answer","").strip(),"correct":True}]
        for j,o in enumerate(options):db.session.add(QuestionOption(question=q,option_text=o["text"].strip(),is_correct=bool(o.get("correct")),option_order=j+1))

def bank_json(item):return {"id":item.id,"type":item.question_type,"text":item.question_text,"marks":item.marks,"topic":item.topic,"difficulty":item.difficulty,"tags":item.tags or [],"options":item.options or [],"correct_answer":item.correct_answer,"created_at":item.created_at.isoformat() if item.created_at else None}

@api.route("/admin/question-bank",methods=["GET","POST"])
@role_required("ADMIN")
def question_bank():
    if request.method=="GET":
        query=QuestionBankItem.query.filter_by(institute_id=tenant_id())
        if request.args.get("topic"):query=query.filter_by(topic=request.args["topic"])
        if request.args.get("difficulty"):query=query.filter_by(difficulty=request.args["difficulty"].upper())
        if request.args.get("q"):query=query.filter(QuestionBankItem.question_text.ilike(f"%{request.args['q']}%"))
        return ok([bank_json(item) for item in query.order_by(QuestionBankItem.created_at.desc()).all()])
    body=request.get_json() or {};qtype=str(body.get("type","SINGLE")).upper();text=str(body.get("text","")).strip();difficulty=str(body.get("difficulty","MEDIUM")).upper()
    if qtype not in {"SINGLE","MULTIPLE","TRUE_FALSE","SHORT_ANSWER","PARAGRAPH"} or not text:return fail("Valid question type and text are required",422)
    if difficulty not in {"EASY","MEDIUM","HARD"}:return fail("Difficulty must be EASY, MEDIUM, or HARD",422)
    item=QuestionBankItem(institute_id=tenant_id(),question_type=qtype,question_text=text,marks=float(body.get("marks",1)),topic=str(body.get("topic","")).strip() or None,difficulty=difficulty,tags=body.get("tags") or [],options=body.get("options") or [],correct_answer=body.get("correct_answer"),created_by=uid());db.session.add(item);db.session.flush();audit("CREATE_BANK_QUESTION",item.id,"QUESTION_BANK");db.session.commit();return ok(bank_json(item),"Question added",201)

@api.route("/admin/question-bank/<int:item_id>",methods=["PUT","DELETE"])
@role_required("ADMIN")
def question_bank_item(item_id):
    item=QuestionBankItem.query.filter_by(id=item_id,institute_id=tenant_id()).first()
    if not item:return fail("Question not found",404)
    if request.method=="DELETE":audit("DELETE_BANK_QUESTION",item.id,"QUESTION_BANK");db.session.delete(item);db.session.commit();return ok(message="Question deleted")
    body=request.get_json() or {}
    for key,attr in (("text","question_text"),("type","question_type"),("marks","marks"),("topic","topic"),("difficulty","difficulty"),("tags","tags"),("options","options"),("correct_answer","correct_answer")):
        if key in body:setattr(item,attr,body[key])
    item.question_type=str(item.question_type).upper();item.difficulty=str(item.difficulty).upper();audit("UPDATE_BANK_QUESTION",item.id,"QUESTION_BANK");db.session.commit();return ok(bank_json(item),"Question updated")

BANK_IMPORT_HEADERS=["Question Type","Question","Marks","Topic","Difficulty","Tags","Options","Correct Answer"]
def parse_bank_rows(upload):
    if upload.filename.lower().endswith(".csv"):rows=list(csv.reader(StringIO(upload.read().decode("utf-8-sig"))))
    elif upload.filename.lower().endswith(".xlsx"):rows=[[cell for cell in row] for row in load_workbook(upload,read_only=True,data_only=True).active.iter_rows(values_only=True)]
    else:raise ValueError("Use a CSV or XLSX file")
    if not rows or [str(x or "").strip() for x in rows[0]][:8]!=BANK_IMPORT_HEADERS:raise ValueError("Question-bank columns do not match the required template")
    parsed=[];errors=[]
    for number,row in enumerate(rows[1:],2):
        if not any(value not in (None,"") for value in row):continue
        try:
            qtype=str(row[0]).strip().upper();text=str(row[1]).strip();difficulty=str(row[4] or "MEDIUM").strip().upper();marks=float(row[2] or 1)
            if qtype not in {"SINGLE","MULTIPLE","TRUE_FALSE","SHORT_ANSWER","PARAGRAPH"} or not text or difficulty not in {"EASY","MEDIUM","HARD"}:raise ValueError("invalid type, question, or difficulty")
            options=[value.strip() for value in str(row[6] or "").split("|") if value.strip()];correct=str(row[7] or "").strip();parsed.append({"type":qtype,"text":text,"marks":marks,"topic":str(row[3] or "").strip(),"difficulty":difficulty,"tags":[x.strip() for x in str(row[5] or "").split(",") if x.strip()],"options":[{"text":x,"correct":x==correct} for x in options],"correct_answer":correct})
        except (ValueError,TypeError) as error:errors.append({"row":number,"error":str(error)})
    return parsed,errors

@api.post("/admin/question-bank/import-preview")
@role_required("ADMIN")
def preview_bank_import():
    try:items,errors=parse_bank_rows(request.files.get("file"));return ok({"items":items,"errors":errors,"valid":len(items),"invalid":len(errors)})
    except (ValueError,AttributeError,UnicodeDecodeError) as error:return fail(str(error),422)

@api.post("/admin/question-bank/import")
@role_required("ADMIN")
def import_bank():
    try:
        items,errors=parse_bank_rows(request.files.get("file"))
        if errors:return fail("Fix validation errors before importing",422)
        for data in items:db.session.add(QuestionBankItem(institute_id=tenant_id(),question_type=data["type"],question_text=data["text"],marks=data["marks"],topic=data["topic"] or None,difficulty=data["difficulty"],tags=data["tags"],options=data["options"],correct_answer=data["correct_answer"],created_by=uid()))
        audit("IMPORT_QUESTION_BANK",None,"QUESTION_BANK",{"count":len(items)},"TRANSACTION");db.session.commit();return ok({"created":len(items)},"Questions imported",201)
    except (ValueError,AttributeError,UnicodeDecodeError) as error:return fail(str(error),422)

def replace_question(q, item):
    question_type=item.get("type",q.question_type).upper()
    if question_type not in {"SINGLE","MULTIPLE","TRUE_FALSE","SHORT_ANSWER","PARAGRAPH"}:
        raise ValueError("Unsupported question type")
    q.question_type=question_type
    q.question_text=item.get("text",q.question_text).strip()
    q.marks=float(item.get("marks",q.marks))
    if "order" in item:q.question_order=int(item["order"])
    q.options.clear();db.session.flush()
    options=item.get("options",[])
    if question_type=="TRUE_FALSE":
        correct=str(item.get("correct_answer","true")).lower()=="true"
        options=[{"text":"True","correct":correct},{"text":"False","correct":not correct}]
    elif question_type in ("SHORT_ANSWER","PARAGRAPH"):
        options=[{"text":item.get("correct_answer","").strip(),"correct":True}]
    for index,option in enumerate(options):
        db.session.add(QuestionOption(question=q,option_text=option["text"].strip(),is_correct=bool(option.get("correct")),option_order=index+1))

@api.post("/admin/exams/<int:eid>/questions")
@role_required("ADMIN")
def create_question(eid):
    e=Exam.query.filter_by(id=eid,institute_id=tenant_id(),archived_at=None).first()
    if not e:return fail("Exam not found",404)
    item=request.get_json() or {}
    item["order"]=len(e.questions)+1
    try:add_questions(e,[item])
    except (KeyError,TypeError,ValueError) as error:return fail(str(error),422)
    e.total_marks=sum(q.marks for q in e.questions);audit("ADD_QUESTION",e.id);db.session.commit()
    return ok(exam_json(e,True),"Question added",201)

@api.put("/admin/exams/<int:eid>/questions/<int:qid>")
@role_required("ADMIN")
def update_question(eid,qid):
    e=Exam.query.filter_by(id=eid,institute_id=tenant_id(),archived_at=None).first()
    q=Question.query.filter_by(id=qid,exam_id=eid).first() if e else None
    if not q:return fail("Draft exam question not found",404)
    try:replace_question(q,request.get_json() or {})
    except (KeyError,TypeError,ValueError) as error:return fail(str(error),422)
    e.total_marks=sum(item.marks for item in e.questions);audit("UPDATE_QUESTION",e.id);db.session.commit()
    return ok(exam_json(e,True),"Question updated")

@api.delete("/admin/exams/<int:eid>/questions/<int:qid>")
@role_required("ADMIN")
def delete_question(eid,qid):
    e=Exam.query.filter_by(id=eid,institute_id=tenant_id(),archived_at=None).first()
    q=Question.query.filter_by(id=qid,exam_id=eid).first() if e else None
    if not q:return fail("Draft exam question not found",404)
    remaining=[item for item in e.questions if item.id!=qid]
    db.session.delete(q)
    for index,item in enumerate(remaining):item.question_order=index+1
    e.total_marks=sum(item.marks for item in remaining);audit("DELETE_QUESTION",e.id);db.session.commit()
    return ok(exam_json(e,True),"Question deleted")

@api.put("/admin/exams/<int:eid>/questions/reorder")
@role_required("ADMIN")
def reorder_questions(eid):
    e=Exam.query.filter_by(id=eid,institute_id=tenant_id(),archived_at=None).first()
    if not e:return fail("Exam not found",404)
    order=(request.get_json() or {}).get("question_ids",[])
    if set(order)!={q.id for q in e.questions}:return fail("Question order must contain every question exactly once",422)
    for index,qid in enumerate(order):db.session.get(Question,qid).question_order=index+1
    audit("REORDER_QUESTIONS",e.id);db.session.commit();return ok(exam_json(e,True),"Questions reordered")
def audit(action,eid=None,entity_type="EXAM",details=None,category="AUDIT",severity="INFO"):
    db.session.add(AuditLog(institute_id=tenant_id(),actor_id=uid(),action=action,category=category,severity=severity,entity_type=entity_type,entity_id=eid,details=details or {},ip_address=client_ip(),user_agent=request.user_agent.string[:500],request_id=request.headers.get("X-Request-ID") or str(uuid.uuid4())))
@api.post("/admin/exams/<int:eid>/<action>")
@role_required("ADMIN")
def lifecycle(eid,action):
    e=Exam.query.filter_by(id=eid,institute_id=tenant_id(),archived_at=None).first()
    if not e:return fail("Exam not found",404)
    transitions={"publish":("DRAFT","PUBLISHED"),"complete":("PUBLISHED","COMPLETED"),"reactivate":("COMPLETED","PUBLISHED")}
    if action not in transitions:return fail("Invalid action",404)
    old,new=transitions[action]
    if e.status!=old:return fail(f"Cannot {action} an exam in {e.status} status",409)
    if action=="publish" and not e.questions:return fail("Add at least one question before publishing",422)
    if action=="publish" and not e.student_category_id and not e.assignments:return fail("Assign the exam to a student type or selected students before publishing",422)
    b=request.get_json(silent=True) or {}
    if action=="reactivate":
        if not b.get("start_at") or not b.get("end_at"):return fail("Choose a new start time and end time before reactivating the exam",422)
        try:start_at=dt(b["start_at"]);end_at=dt(b["end_at"])
        except (TypeError,ValueError):return fail("Choose valid start and end times",422)
        start_check=start_at.replace(tzinfo=timezone.utc) if start_at.tzinfo is None else start_at;end_check=end_at.replace(tzinfo=timezone.utc) if end_at.tzinfo is None else end_at
        if end_check<=start_check:return fail("End time must be later than start time",422)
        if end_check<=now():return fail("End time must be in the future",422)
        e.start_at=start_at;e.end_at=end_at;e.completed_at=None;e.student_category_id=None
        active_students=User.query.filter_by(institute_id=tenant_id(),role="STUDENT",status="ACTIVE").all();e.assignments.clear();db.session.flush()
        for student in active_students:e.assignments.append(ExamAssignment(student_id=student.id,assigned_by=uid()))
        for attempt in e.attempts:
            if attempt.status in ("COMPLETED","TERMINATED"):attempt.retake_allowed=True;attempt.retake_reason="New institute-wide exam cycle created by reactivation"
    e.status=new
    if new=="PUBLISHED":e.published_at=now()
    if new=="COMPLETED":e.completed_at=now()
    audit(action.upper()+"_EXAM",e.id);db.session.commit();return ok(exam_json(e),f"Exam {new.lower()}")
@api.get("/admin/history")
@role_required("ADMIN")
def admin_history(): return ok([{"exam":exam_json(e),"attempts":[attempt_json(a) for a in e.attempts]} for e in Exam.query.filter_by(institute_id=tenant_id()).order_by(Exam.created_at.desc()).all()])
@api.get("/admin/attempts/<int:aid>")
@role_required("ADMIN")
def admin_attempt(aid):
    a=ExamAttempt.query.join(Exam).filter(ExamAttempt.id==aid,Exam.institute_id==tenant_id()).first();return ok(attempt_json(a,True)) if a else fail("Attempt not found",404)

def monitored_attempt_json(a):
    t=now();last=a.last_activity_at or a.started_at;last=last.replace(tzinfo=timezone.utc) if last.tzinfo is None else last
    deadline=attempt_deadline(a);answered=len(a.answers);remaining=max(0,int((deadline-t).total_seconds())) if a.status in ("IN_PROGRESS","PAUSED") else 0
    return {**attempt_json(a),"attempt_number":a.attempt_number,"answered":answered,"unanswered":max(0,a.total_questions-answered),"progress":round(100*answered/a.total_questions,1) if a.total_questions else 0,"remaining_seconds":remaining,"last_activity_at":last.isoformat(),"connection_status":"ONLINE" if a.status in ("IN_PROGRESS","PAUSED") and (t-last).total_seconds()<45 else "DISCONNECTED","auto_submitted":a.auto_submitted,"refresh_count":a.refresh_count,"risk_score":a.risk_score,"warning":a.warning_message,"termination_reason":a.termination_reason,"events":[{"id":event.id,"type":event.event_type,"severity":event.severity,"risk_points":event.risk_points,"details":event.details,"ip_address":event.ip_address,"created_at":event.created_at.isoformat()} for event in sorted(a.events,key=lambda item:item.created_at,reverse=True)[:30]]}

@api.get("/admin/monitoring")
@role_required("ADMIN")
def monitoring():
    cache_key=f"monitoring:{tenant_id()}:{request.args.get('status','all')}";cached=get_json(cache_key)
    if cached:return ok(cached)
    query=ExamAttempt.query.options(joinedload(ExamAttempt.exam),joinedload(ExamAttempt.student).joinedload(User.profile),selectinload(ExamAttempt.answers),selectinload(ExamAttempt.events)).join(Exam).filter(Exam.institute_id==tenant_id())
    status=request.args.get("status")
    if status:query=query.filter(ExamAttempt.status==status.upper())
    attempts=query.order_by(ExamAttempt.last_activity_at.desc()).limit(200).all();items=[monitored_attempt_json(a) for a in attempts]
    data={"attempts":items,"summary":{"active":sum(item["status"] in ("IN_PROGRESS","PAUSED") for item in items),"disconnected":sum(item["connection_status"]=="DISCONNECTED" and item["status"] in ("IN_PROGRESS","PAUSED") for item in items),"high_risk":sum(item["risk_score"]>=20 for item in items),"auto_submitted":sum(bool(item["auto_submitted"]) for item in items)}};set_json(cache_key,data,current_app.config["MONITORING_CACHE_SECONDS"]);return ok(data)

@api.post("/admin/attempts/<int:aid>/control")
@role_required("ADMIN")
def control_attempt(aid):
    a=ExamAttempt.query.join(Exam).filter(ExamAttempt.id==aid,Exam.institute_id==tenant_id()).first()
    if not a:return fail("Attempt not found",404)
    body=request.get_json() or {};action=str(body.get("action","")).upper();before={"status":a.status,"extra_time_minutes":a.extra_time_minutes,"retake_allowed":a.retake_allowed}
    if action=="ADD_TIME":
        minutes=int(body.get("minutes",0))
        if minutes<1 or minutes>180:return fail("Extra time must be between 1 and 180 minutes",422)
        a.extra_time_minutes+=minutes
    elif action=="PAUSE":
        if a.status!="IN_PROGRESS":return fail("Only an active attempt can be paused",409)
        a.status="PAUSED";a.paused_at=now()
    elif action=="RESUME":
        if a.status!="PAUSED":return fail("Only a paused attempt can be resumed",409)
        paused=a.paused_at.replace(tzinfo=timezone.utc) if a.paused_at and a.paused_at.tzinfo is None else a.paused_at
        a.paused_seconds+=(int((now()-paused).total_seconds()) if paused else 0);a.paused_at=None;a.status="IN_PROGRESS"
    elif action=="TERMINATE":
        if a.status not in ("IN_PROGRESS","PAUSED"):return fail("Attempt is not active",409)
        a.status="TERMINATED";a.termination_reason=str(body.get("reason") or "Terminated by administrator")[:500]
    elif action=="FORCE_SUBMIT":
        if a.status not in ("IN_PROGRESS","PAUSED"):return fail("Attempt is not active",409)
        submit_attempt(a)
    elif action in ("REOPEN","ALLOW_RETAKE"):
        if a.status not in ("COMPLETED","TERMINATED"):return fail("Only a finished attempt can be reopened",409)
        reason=str(body.get("reason","")).strip()
        if not reason:return fail("A reason is required",422)
        a.retake_allowed=True;a.retake_reason=reason[:500]
        exam=a.exam;current=now();exam_end=exam.end_at.replace(tzinfo=timezone.utc) if exam.end_at and exam.end_at.tzinfo is None else exam.end_at
        if exam.status!="PUBLISHED" or not exam_end or exam_end<=current:
            exam.status="PUBLISHED";exam.start_at=current;exam.end_at=current+timedelta(days=7);exam.published_at=current;exam.completed_at=None
        if not any(item.student_id==a.student_id for item in exam.assignments):exam.assignments.append(ExamAssignment(student_id=a.student_id,assigned_by=uid()))
    elif action=="WARN":
        message=str(body.get("message","")).strip()
        if not message:return fail("Warning message is required",422)
        a.warning_message=message[:500];a.warning_sent_at=now()
    else:return fail("Unsupported attempt control",422)
    after={"status":a.status,"extra_time_minutes":a.extra_time_minutes,"retake_allowed":a.retake_allowed}
    audit(action,a.id,"ATTEMPT",{"before":before,"after":after,"reason":body.get("reason"),"message":body.get("message")},"SECURITY","WARNING" if action in ("TERMINATE","WARN") else "INFO")
    db.session.commit();return ok(monitored_attempt_json(a),"Attempt updated")

@api.get("/admin/audit-logs")
@role_required("ADMIN")
def audit_logs():
    query=AuditLog.query.filter_by(institute_id=tenant_id())
    if request.args.get("category"):query=query.filter_by(category=request.args["category"].upper())
    if request.args.get("action"):query=query.filter(AuditLog.action.ilike(f"%{request.args['action']}%"))
    if request.args.get("severity"):query=query.filter_by(severity=request.args["severity"].upper())
    if request.args.get("from"):query=query.filter(AuditLog.created_at>=dt(request.args["from"]))
    if request.args.get("to"):query=query.filter(AuditLog.created_at<=dt(request.args["to"]))
    page=max(1,int(request.args.get("page",1)));per_page=min(max(1,int(request.args.get("per_page",50))),200);total=query.count();rows=query.order_by(AuditLog.created_at.desc()).offset((page-1)*per_page).limit(per_page).all()
    actors={user.id:user.name for user in User.query.filter(User.id.in_({row.actor_id for row in rows if row.actor_id})).all()}
    items=[{"id":row.id,"action":row.action,"category":row.category,"severity":row.severity,"entity_type":row.entity_type,"entity_id":row.entity_id,"details":row.details,"ip_address":row.ip_address,"user_agent":row.user_agent,"request_id":row.request_id,"created_at":row.created_at.isoformat(),"actor":actors.get(row.actor_id,"System")} for row in rows]
    return ok({"items":items,"pagination":{"page":page,"per_page":per_page,"total":total,"pages":max(1,(total+per_page-1)//per_page)}})

@api.get("/admin/audit-logs/export")
@role_required("ADMIN")
def export_audit_logs():
    rows=AuditLog.query.filter_by(institute_id=tenant_id()).order_by(AuditLog.created_at.desc()).limit(10000).all();output=StringIO();writer=csv.writer(output);writer.writerow(["Timestamp","Category","Severity","Actor","Action","Entity","Entity ID","IP Address","Request ID"])
    for row in rows:writer.writerow([row.created_at.isoformat(),row.category,row.severity,db.session.get(User,row.actor_id).name if row.actor_id else "System",row.action,row.entity_type,row.entity_id,row.ip_address,row.request_id])
    audit("EXPORT_AUDIT_LOGS",None,"AUDIT_LOG",{"rows":len(rows)},"TRANSACTION");db.session.commit();return current_app.response_class(output.getvalue(),mimetype="text/csv",headers={"Content-Disposition":"attachment; filename=audit-logs.csv"})
